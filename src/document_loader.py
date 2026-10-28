from pathlib import Path

from pypdf import PdfReader
from docx import Document


def load_document(file_path: str) -> str:

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = path.suffix.lower()

    if extension in (".txt", ".md"):

        return path.read_text(
            encoding="utf-8"
        )

    elif extension == ".pdf":

        # pypdf splits many words of the EU Official Journal PDFs into
        # letter groups ("r ights", "Uni on", "REGUL A TION").  That corrupts
        # the input of every method and especially breaks the exact-copy
        # validation of extract-support.  pypdfium2 extracts these PDFs
        # cleanly, so it is used first; pypdf stays as a fallback.
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            pages = []
            for page in pdf:
                text = page.get_textpage().get_text_range()
                text = text.replace("\r\n", "\n").replace("\r", "\n")
                if text.strip():
                    pages.append(text)
            pdf.close()
            return "\n\n".join(pages)

        except ImportError:
            pass

        reader = PdfReader(
            str(path)
        )

        pages = []

        for page in reader.pages:

            text = page.extract_text()

            if text:
                pages.append(text)

        return "\n\n".join(pages)

    elif extension in (".html", ".htm"):

        return html_to_markdown(path.read_text(encoding="utf-8", errors="replace"))

    elif extension == ".docx":

        document = Document(
            str(path)
        )

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return "\n\n".join(
            paragraphs
        )

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )


def html_to_markdown(html: str) -> str:
    """
    HTML (e.g. the EUR-Lex HTML version of an act) -> Markdown text.
    Avoids PDF extraction problems (split words, headers, footers); keeps
    headings, lists and tables.  Uses markdownify, BeautifulSoup as fallback.
    """
    import re

    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "noscript", "head", "nav", "footer"]):
            tag.decompose()
        html = str(soup)
    except ImportError:
        soup = None

    try:
        from markdownify import markdownify
        text = markdownify(html, heading_style="ATX", strip=["a", "img"])
    except ImportError:
        if soup is None:
            raise ImportError("install markdownify or beautifulsoup4 to read HTML files")
        text = soup.get_text("\n")

    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
